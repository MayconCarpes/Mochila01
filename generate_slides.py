import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    DARK_BLUE = RGBColor(15, 32, 67)
    ACCENT_BLUE = RGBColor(0, 120, 215)
    LIGHT_BG = RGBColor(245, 247, 250)
    DARK_TEXT = RGBColor(30, 30, 30)
    WHITE = RGBColor(255, 255, 255)
    GREEN_TEXT = RGBColor(16, 124, 65)

    def add_header(slide, title_text, category_text="PROBLEMA DA MOCHILA 0/1"):
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = DARK_BLUE
        top_bar.line.color.rgb = DARK_BLUE

        tf = top_bar.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.8)
        tf.margin_top = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_BLUE

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

    # SLIDE 1: Cover
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BLUE
    bg1.line.fill.background()

    txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
    tf1 = txBox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "PROBLEMA DA MOCHILA 0/1"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf1.add_paragraph()
    p.text = "Resolução por Meta-Heurísticas Local e Populacional"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p = tf1.add_paragraph()
    p.text = "Análise Comparativa: Busca Tabu vs Algoritmo Genético em 9 Instâncias Benchmark"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(200, 215, 240)

    p = tf1.add_paragraph()
    p.text = "\nAutor: Maycon Carpes"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # SLIDE 2: Definição do Problema
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide2)
    add_header(slide2, "Definição e Formulação Matemática do Problema")

    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = WHITE
    card1.line.color.rgb = RGBColor(220, 225, 230)
    
    tf = card1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = " Conceito Geral"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    bullets = [
        ("Definição: ", "Dado um conjunto de N itens, cada um com peso w_i e lucro p_i, e uma mochila com capacidade C."),
        ("Objetivo: ", "Selecionar um subconjunto de itens que maximize o lucro total sem que o peso exceda a capacidade."),
        ("Natureza Combinatória: ", "O número de soluções possíveis é 2^N. Trata-se de um problema NP-Difícil (NP-Hard)."),
        ("Necessidade de Heurísticas: ", "Para N elevado (ex: N=100), a busca exaustiva se torna inviável, justificando o uso de Meta-heurísticas.")
    ]
    for title, desc in bullets:
        p = tf.add_paragraph()
        p.text = f"• {title}"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_TEXT
        run = p.add_run()
        run.text = desc
        run.font.bold = False

    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    card2.fill.solid()
    card2.fill.fore_color.rgb = WHITE
    card2.line.color.rgb = RGBColor(220, 225, 230)

    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_right = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = " Formulação Matemática"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    math_points = [
        ("Variável de Decisão:", "\n  x_i ∈ {0, 1} para cada item i=1..N\n  (1 se o item for incluído, 0 caso contrário)"),
        ("Função Objetivo (Maximizar Lucro):", "\n  Maximize  Z = ∑ (p_i * x_i)  para i=1..N"),
        ("Restrição de Capacidade:", "\n  Sujeito a  ∑ (w_i * x_i) ≤ C  para i=1..N"),
        ("Desafio da Factibilidade:", "\n  Soluções onde o peso excede C são infactíveis e precisam de estratégias de reparo.")
    ]
    for title, desc in math_points:
        p = tf2.add_paragraph()
        p.text = f"• {title}"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_TEXT
        run = p.add_run()
        run.text = desc
        run.font.bold = False

    # SLIDE 3: Instâncias Benchmark
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide3)
    add_header(slide3, "Conjunto de Dados — 9 Instâncias Benchmark")

    rows, cols = 10, 5
    table_shape = slide3.shapes.add_table(rows, cols, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.4))
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(2.5)
    table.columns[3].width = Inches(2.5)
    table.columns[4].width = Inches(2.733)

    headers = ["Instância", "Nº de Itens (N)", "Capacidade (C)", "Lucro Ótimo", "Fonte do Dataset"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    data_inst = [
        ("P01", "10", "165", "309", "Burkardt (FSU)"),
        ("P02", "5", "26", "51", "Burkardt (FSU)"),
        ("P03", "6", "190", "150", "Burkardt (FSU)"),
        ("P04", "7", "50", "107", "Burkardt (FSU)"),
        ("P05", "8", "104", "900", "Burkardt (FSU)"),
        ("P06", "7", "170", "1.735", "Burkardt (FSU)"),
        ("P07", "15", "750", "1.458", "Burkardt (FSU)"),
        ("P08", "24", "6.404.180", "13.549.094", "Burkardt (FSU)"),
        ("knapPI_1_100", "100", "995", "9.147", "Pisinger (Large Scale)")
    ]

    for i, row in enumerate(data_inst):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(238, 242, 248)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(12)
            p.font.color.rgb = DARK_TEXT
            p.alignment = PP_ALIGN.CENTER
            if j == 0:
                p.font.bold = True

    # SLIDE 4: Arquitetura
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide4)
    add_header(slide4, "Arquitetura do Código e Organização de Módulos")

    card_arch = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    card_arch.fill.solid()
    card_arch.fill.fore_color.rgb = WHITE
    card_arch.line.color.rgb = RGBColor(220, 225, 230)

    tf = card_arch.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = " Organização dos Módulos (Clean Code & Princípio da Responsabilidade Única)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    modules = [
        ("src/models/", "Definições das Estruturas de Dados via dataclass (KnapsackInstance, ExecutionResult)."),
        ("src/utils/", "Heurísticas auxiliares: greedy_solution (ordenação por lucro/peso) e repair_solution (reparo de factibilidade)."),
        ("src/algorithms/", "Implementação polimórfica herdando de BaseAlgorithm (TabuSearch e GeneticAlgorithm)."),
        ("src/runners/", "Orquestrador de experimentos (ExperimentRunner) e Gerador automático de relatórios CSV e Word DOCX (ReportGenerator)."),
        ("data/instances.py", "Módulo centralizador contendo os dados das 9 instâncias benchmark pré-configuradas."),
        ("main.py", "Ponto de entrada do sistema: constrói a grade de hiperparâmetros, executa os experimentos e compila as métricas.")
    ]
    for mod, desc in modules:
        p = tf.add_paragraph()
        p.text = f"• {mod}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = ACCENT_BLUE
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_TEXT

    # SLIDE 5: Busca Tabu - Conceitos
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide5)
    add_header(slide5, "Meta-heurística Local — Busca Tabu")

    c1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = " Visão Geral da Busca Tabu"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    bt_points = [
        ("Tipo de Busca: ", "Meta-heurística baseada em solução única (busca local em vizinhança)."),
        ("Inicialização: ", "Solução inicial gerada pela Heurística Gulosa por razão lucro/peso."),
        ("Estrutura de Memória: ", "Lista Tabu FIFO de tamanho fixo para evitar ciclagem e escapar de mínimos locais."),
        ("Definição de Vizinhança: ", "Movimento de Flip de 1-bit (inverter inclusão/exclusão de um item por vez).")
    ]
    for t, d in bt_points:
        p = tf.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    c2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = " Mecanismos de Controle"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    bt_mech = [
        ("Critério de Aspiração: ", "Um movimento tabu PODE ser aceito se resultar em um lucro estritamente SUPERIOR à melhor solução global encontrada até então."),
        ("Reparo de Vizinhança: ", "Cada vizinho gerado passa pelo repair_solution para remover itens de menor razão lucro/peso se exceder a capacidade C."),
        ("Critério de Parada: ", "Número máximo de iterações atingido ou ausência de vizinhos válidos.")
    ]
    for t, d in bt_mech:
        p = tf2.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    # SLIDE 6: Busca Tabu - Hiperparâmetros
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide6)
    add_header(slide6, "Busca Tabu — Hiperparâmetros Testados")

    t_shape = slide6.shapes.add_table(4, 3, Inches(1.5), Inches(2.0), Inches(10.333), Inches(4.2))
    table = t_shape.table
    table.columns[0].width = Inches(3.4)
    table.columns[1].width = Inches(3.4)
    table.columns[2].width = Inches(3.533)

    headers = ["Tamanho da Lista Tabu", "Iterações Máximas", "Total de Combinações / Testes"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    bt_params = [
        ("5", "100", "9 configurações por instância"),
        ("10", "500", "x 9 instâncias benchmark"),
        ("20", "1.000", "= 81 cenários x 3 repetições = 243 execuções")
    ]
    for i, row in enumerate(bt_params):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(238, 242, 248)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(13)
            p.alignment = PP_ALIGN.CENTER
            if j == 2:
                p.font.bold = True

    # SLIDE 7: Algoritmo Genético - Conceitos
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide7)
    add_header(slide7, "Meta-heurística Populacional — Algoritmo Genético")

    c1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = " Representação e Operadores Genéticos"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    ga_points = [
        ("Cromossomo: ", "Vetor binário de tamanho N, onde x_i = 1 indica inclusão do item i."),
        ("Inicialização: ", "População aleatória combinada com 1 indivíduo da Heurística Gulosa (seeding)."),
        ("Seleção por Torneio: ", "Torneio k=3 para selecionar os genitores com maior fitness (lucro)."),
        ("Crossover: ", "Crossover de 1 ponto com taxa configurável (80%).")
    ]
    for t, d in ga_points:
        p = tf.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    c2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = " Diversidade e Manutenção de Elitismo"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    ga_mech = [
        ("Mutação Bit-Flip: ", "Inversão binária por gene com taxas de 1% e 5% para manutenção da variabilidade genética."),
        ("Elitismo Estrito: ", "O melhor indivíduo da geração é preservado inalterado diretamente na próxima população."),
        ("Garantia de Factibilidade: ", "Todos os descendentes mutados passam obrigatoriamente pelo reparo guloso de capacidade.")
    ]
    for t, d in ga_mech:
        p = tf2.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    # SLIDE 8: Algoritmo Genético - Hiperparâmetros
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide8)
    add_header(slide8, "Algoritmo Genético — Hiperparâmetros Testados")

    t_shape = slide8.shapes.add_table(4, 4, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.2))
    table = t_shape.table
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(2.8)
    table.columns[3].width = Inches(3.333)

    headers = ["Tamanho da População", "Número de Gerações", "Taxa de Mutação", "Total de Execuções"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    ga_params = [
        ("50", "100", "0.01 (1%)", "12 configurações por instância"),
        ("100", "500", "0.05 (5%)", "x 9 instâncias benchmark"),
        ("Crossover: 0.8", "1.000", "Elitismo: 1 ind.", "= 108 cenários x 3 rep. = 324 exec.")
    ]
    for i, row in enumerate(ga_params):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(238, 242, 248)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(13)
            p.alignment = PP_ALIGN.CENTER
            if j == 3:
                p.font.bold = True

    # SLIDE 9: Tabela Comparativa de Gap (%)
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide9)
    add_header(slide9, "Resultados Comparativos — Qualidade da Solução e Gap (%)")

    table_shape = slide9.shapes.add_table(10, 7, Inches(0.5), Inches(1.3), Inches(12.333), Inches(5.6))
    table = table_shape.table

    col_widths = [Inches(2.2), Inches(1.1), Inches(2.0), Inches(2.0), Inches(1.6), Inches(1.9), Inches(1.533)]
    for j, w in enumerate(col_widths):
        table.columns[j].width = w

    headers = ["Instância", "Itens", "Ótimo Conhecido", "Busca Tabu (Lucro)", "Gap BT (%)", "Alg. Genético (Lucro)", "Gap AG (%)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    res_data = [
        ("P01", "10", "309", "309", "0,00%", "309", "0,00%"),
        ("P02", "5", "51", "47", "7,84%", "51", "0,00%"),
        ("P03", "6", "150", "150", "0,00%", "150", "0,00%"),
        ("P04", "7", "107", "107", "0,00%", "107", "0,00%"),
        ("P05", "8", "900", "858", "4,67%", "900", "0,00%"),
        ("P06", "7", "1.735", "1.735", "0,00%", "1.735", "0,00%"),
        ("P07", "15", "1.458", "1.456", "0,14%", "1.458", "0,00%"),
        ("P08", "24", "13.549.094", "13.449.995", "0,73%", "13.549.094", "0,00%"),
        ("knapPI_1_100", "100", "9.147", "8.817", "3,61%", "9.147", "0,00%")
    ]

    for i, row in enumerate(res_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(238, 242, 248)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.alignment = PP_ALIGN.CENTER
            if j in [0, 2]:
                p.font.bold = True
            if j == 4:
                p.font.color.rgb = GREEN_TEXT if val == "0,00%" else RGBColor(180, 40, 40)
            if j == 6:
                p.font.bold = True
                p.font.color.rgb = GREEN_TEXT

    # SLIDE 10: Tempos de Execução e Convergência
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide10)
    add_header(slide10, "Análise de Desempenho — Tempo de Execução")

    table_shape = slide10.shapes.add_table(10, 6, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.4))
    table = table_shape.table

    widths = [Inches(2.5), Inches(1.5), Inches(2.2), Inches(2.2), Inches(1.8), Inches(1.533)]
    for j, w in enumerate(widths):
        table.columns[j].width = w

    headers = ["Instância", "Itens (N)", "Tempo Busca Tabu (s)", "Tempo Alg. Genético (s)", "Gap BT (%)", "Gap AG (%)"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    perf_data = [
        ("P01", "10", "0,0017 s", "0,0209 s", "0,00%", "0,00%"),
        ("P02", "5", "0,0000 s", "0,0164 s", "7,84%", "0,00%"),
        ("P03", "6", "0,0007 s", "0,0168 s", "0,00%", "0,00%"),
        ("P04", "7", "0,0009 s", "0,0174 s", "0,00%", "0,00%"),
        ("P05", "8", "0,0011 s", "0,0180 s", "4,67%", "0,00%"),
        ("P06", "7", "0,0010 s", "0,0178 s", "0,00%", "0,00%"),
        ("P07", "15", "0,0034 s", "0,0216 s", "0,14%", "0,00%"),
        ("P08", "24", "0,0071 s", "0,0324 s", "0,73%", "0,00%"),
        ("knapPI_1_100", "100", "0,1039 s", "0,0838 s", "3,61%", "0,00%")
    ]

    for i, row in enumerate(perf_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(238, 242, 248)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.alignment = PP_ALIGN.CENTER
            if j == 0:
                p.font.bold = True

    # SLIDE 11: Conclusões
    slide11 = prs.slides.add_slide(blank_slide_layout)
    set_bg(slide11)
    add_header(slide11, "Conclusões e Discussão dos Resultados")

    c1 = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    tf = c1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = " Qualidade da Solução"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    conc_q = [
        ("Algoritmo Genético: ", "Alcançou 100% de sucesso nas 9 instâncias (Gap Médio de 0,00%)."),
        ("Busca Tabu: ", "Atingiu o ótimo em 4 de 9 instâncias (Gap Médio de 1,89%)."),
        ("Escalabilidade (100 itens): ", "O AG manteve o ótimo em N=100 (knapPI_1_100), enquanto a Busca Tabu apresentou gap de 3,61%.")
    ]
    for t, d in conc_q:
        p = tf.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    c2 = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = " Eficiência e Trade-off"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE

    conc_t = [
        ("Velocidade da Busca Tabu: ", "Extremamente rápida para instâncias pequenas (milissegundos), mas com tendência a travar em ótimos locais."),
        ("Poder da População no AG: ", "A diversidade genética aliada ao elitismo impediu a convergência prematura nas instâncias maiores."),
        ("Recomendação Final: ", "O Algoritmo Genético se mostrou superior em qualidade global de otimização para o Problema da Mochila 0/1.")
    ]
    for t, d in conc_t:
        p = tf2.add_paragraph()
        p.text = f"• {t}"
        p.font.bold = True
        p.font.size = Pt(13)
        run = p.add_run()
        run.text = d
        run.font.bold = False

    # SLIDE 12: End / Questions
    slide12 = prs.slides.add_slide(blank_slide_layout)
    bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = DARK_BLUE
    bg12.line.fill.background()

    txBox = slide12.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.333), Inches(3.0))
    tf12 = txBox.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "Obrigado!"
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = WHITE

    p = tf12.add_paragraph()
    p.text = "Perguntas & Discussão"
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = ACCENT_BLUE

    p = tf12.add_paragraph()
    p.text = "\nRepositório: github.com/MayconCarpes/Mochila01"
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(200, 215, 240)

    output_path = "Mochila01_MetaHeuristicas_TabuVsAG.pptx"
    prs.save(output_path)
    print(f"Apresentação gerada com sucesso em '{output_path}'!")

if __name__ == "__main__":
    create_presentation()
