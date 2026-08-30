import csv
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT


class ReportGenerator:
    def __init__(self, results):
        self.results = results

    def generate_csv(self, filepath):
        headers = [
            "Instancia", "N_Itens", "Algoritmo", "Parametros",
            "Melhor_Lucro", "Otimo_Conhecido", "Gap(%)",
            "Tempo_Medio(s)", "Iteracoes/Geracoes",
        ]
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(headers)
            for r in self.results:
                params_str = ", ".join(f"{k}={v}" for k, v in r.params.items())
                writer.writerow([
                    r.instance_name,
                    len(r.best_solution),
                    r.algorithm_name,
                    params_str,
                    r.best_profit,
                    r.optimal_profit,
                    r.gap_percent,
                    r.time_seconds,
                    r.iterations,
                ])

    def generate_docx(self, filepath):
        doc = Document()
        self._add_title(doc)
        self._add_introduction(doc)
        self._add_algorithm_descriptions(doc)
        self._add_results_section(doc)
        self._add_comparative_analysis(doc)
        self._add_conclusion(doc)
        doc.save(filepath)

    def _add_title(self, doc):
        doc.add_heading(
            "Problema da Mochila 0/1 — Relatorio de Comparacao de Meta-heuristicas",
            level=0,
        )

    def _add_introduction(self, doc):
        doc.add_heading("1. Introducao", level=1)
        doc.add_paragraph(
            "Este relatorio apresenta a comparacao entre duas meta-heuristicas aplicadas ao "
            "Problema da Mochila 0/1 (0/1 Knapsack Problem): Busca Tabu (meta-heuristica local) "
            "e Algoritmo Genetico (meta-heuristica populacional). Os algoritmos foram testados em "
            "9 instancias benchmark, variando de 5 a 100 itens."
        )

    def _add_algorithm_descriptions(self, doc):
        doc.add_heading("2. Descricao dos Algoritmos", level=1)

        doc.add_heading("2.1 Busca Tabu", level=2)
        doc.add_paragraph(
            "A Busca Tabu e uma meta-heuristica de busca local que utiliza uma estrutura de memoria "
            "(lista tabu) para evitar ciclos e permitir a exploracao de regioes nao visitadas do espaco "
            "de solucoes. A vizinhanca e definida pelo flip de um bit (adicionar/remover um item). "
            "O criterio de aspiracao permite aceitar movimentos tabu quando geram uma solucao melhor "
            "que a melhor global. A solucao inicial e gerada por uma heuristica gulosa baseada na "
            "razao lucro/peso."
        )
        doc.add_paragraph("Parametros testados:")
        doc.add_paragraph("Tamanho da lista tabu: 5, 10, 20", style="List Bullet")
        doc.add_paragraph("Numero maximo de iteracoes: 100, 500, 1000", style="List Bullet")

        doc.add_heading("2.2 Algoritmo Genetico", level=2)
        doc.add_paragraph(
            "O Algoritmo Genetico e uma meta-heuristica populacional inspirada na evolucao biologica. "
            "Utiliza representacao binaria natural para o problema da mochila 0/1. A selecao e feita "
            "por torneio de tamanho 3, o crossover e de um ponto, e a mutacao e por bit-flip. "
            "O elitismo preserva o melhor individuo entre geracoes. Solucoes infactiveis sao reparadas "
            "removendo itens de menor razao lucro/peso."
        )
        doc.add_paragraph("Parametros testados:")
        doc.add_paragraph("Tamanho da populacao: 50, 100", style="List Bullet")
        doc.add_paragraph("Numero de geracoes: 100, 500, 1000", style="List Bullet")
        doc.add_paragraph("Taxa de mutacao: 0.01, 0.05", style="List Bullet")
        doc.add_paragraph("Taxa de crossover: 0.8", style="List Bullet")

    def _add_hyperlink(self, paragraph, text, url):
        part = paragraph.part
        r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)

        run_elem = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        rPr.append(color)

        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rPr.append(underline)

        run_elem.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        run_elem.append(t)

        hyperlink.append(run_elem)
        paragraph._p.append(hyperlink)

    def _add_results_section(self, doc):
        doc.add_heading("3. Resultados", level=1)

        p = doc.add_paragraph("Os resultados completos estao disponiveis na planilha: ")
        self._add_hyperlink(p, "resultados.csv", "resultados.csv")
        doc.add_paragraph("")

        instances = {}
        for r in self.results:
            instances.setdefault(r.instance_name, []).append(r)

        for inst_name, inst_results in instances.items():
            doc.add_heading(f"Instancia {inst_name}", level=2)
            self._add_results_table(doc, inst_results)
            doc.add_paragraph("")

    def _add_results_table(self, doc, inst_results):
        headers = [
            "Algoritmo", "Parametros", "Melhor Lucro",
            "Otimo", "Gap(%)", "Tempo(s)", "Iter/Ger",
        ]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)

        for r in inst_results:
            row = table.add_row()
            params_str = ", ".join(f"{k}={v}" for k, v in r.params.items())
            values = [
                r.algorithm_name, params_str, str(r.best_profit),
                str(r.optimal_profit), str(r.gap_percent),
                str(r.time_seconds), str(r.iterations),
            ]
            for i, v in enumerate(values):
                row.cells[i].text = v
                for run in row.cells[i].paragraphs[0].runs:
                    run.font.size = Pt(8)

    def _add_comparative_analysis(self, doc):
        doc.add_heading("4. Analise Comparativa", level=1)

        tabu_results = [r for r in self.results if "Tabu" in r.algorithm_name]
        ga_results = [r for r in self.results if "Genetico" in r.algorithm_name]

        if not tabu_results or not ga_results:
            return

        tabu_best = self._best_by_instance(tabu_results)
        ga_best = self._best_by_instance(ga_results)

        doc.add_heading("4.1 Qualidade da Solucao", level=2)
        self._add_comparison_table(doc, tabu_best, ga_best)

        doc.add_heading("4.2 Tempo de Execucao", level=2)
        doc.add_paragraph(
            "A Busca Tabu tende a ser mais rapida em instancias pequenas devido a vizinhanca "
            "simples (flip de 1 bit). O Algoritmo Genetico possui overhead de gerenciamento "
            "populacional, mas escala melhor para instancias maiores pela diversidade de busca."
        )

        doc.add_heading("4.3 Convergencia", level=2)
        doc.add_paragraph(
            "A Busca Tabu converge rapidamente para boas solucoes, especialmente com a "
            "inicializacao gulosa. O Algoritmo Genetico apresenta convergencia mais gradual, "
            "mas com potencial de escapar de otimos locais gracas a diversidade populacional."
        )

    def _best_by_instance(self, results):
        best = {}
        for r in results:
            if r.instance_name not in best or r.best_profit > best[r.instance_name].best_profit:
                best[r.instance_name] = r
        return best

    def _add_comparison_table(self, doc, tabu_best, ga_best):
        headers = [
            "Instancia", "Otimo", "Tabu (Melhor)",
            "Gap Tabu(%)", "AG (Melhor)", "Gap AG(%)",
        ]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)

        all_instances = sorted(set(list(tabu_best.keys()) + list(ga_best.keys())))
        for inst in all_instances:
            row = table.add_row()
            t = tabu_best.get(inst)
            g = ga_best.get(inst)
            values = [
                inst,
                str(t.optimal_profit if t else ""),
                str(t.best_profit if t else ""),
                str(t.gap_percent if t else ""),
                str(g.best_profit if g else ""),
                str(g.gap_percent if g else ""),
            ]
            for i, v in enumerate(values):
                row.cells[i].text = v
                for run in row.cells[i].paragraphs[0].runs:
                    run.font.size = Pt(8)

    def _add_conclusion(self, doc):
        doc.add_heading("5. Conclusao", level=1)

        tabu_results = [r for r in self.results if "Tabu" in r.algorithm_name]
        ga_results = [r for r in self.results if "Genetico" in r.algorithm_name]

        tabu_wins = ga_wins = ties = 0
        for inst in set(r.instance_name for r in self.results):
            tabu_max = max(
                (r.best_profit for r in tabu_results if r.instance_name == inst),
                default=0,
            )
            ga_max = max(
                (r.best_profit for r in ga_results if r.instance_name == inst),
                default=0,
            )
            if tabu_max > ga_max:
                tabu_wins += 1
            elif ga_max > tabu_max:
                ga_wins += 1
            else:
                ties += 1

        doc.add_paragraph(
            f"Nos testes realizados, a Busca Tabu obteve melhor resultado em {tabu_wins} instancia(s), "
            f"o Algoritmo Genetico em {ga_wins} instancia(s), e houve empate em {ties} instancia(s). "
            "Ambas as meta-heuristicas sao eficazes para instancias pequenas a medias do problema "
            "da mochila 0/1, com trade-offs entre velocidade de convergencia (Busca Tabu) e "
            "diversidade de exploracao (Algoritmo Genetico)."
        )
